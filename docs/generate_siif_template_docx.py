# -*- coding: utf-8 -*-
"""
Genera documentos Word (.docx) a partir de la plantilla institucional MinTrabajo
(Arquitectura_Interoperabilidad_Mintrabajo_NewFormat.docx) y los Markdown en este directorio.

Requisitos: py -3 con python-docx y matplotlib.

Ejecución (desde la carpeta docs o la raíz del repo):

    py -3 docs/generate_siif_template_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent

# Plantilla proporcionada (ruta absoluta solicitada por negocio).
TEMPLATE_DOCX = Path(
    r"d:\Ssf\Requerimientos\Desarrollo Interoperbilidad MinTrabajo\Documentación tecnica"
    r"\Arquitectura_Interoperabilidad_Mintrabajo_NewFormat.docx"
)

# La plantilla incluye portada y bloque inicial en párrafos 0..95; el contenido técnico
# del documento MinTrabajo original comienza en el párrafo 96 ("OBJETIVO DEL DOCUMENTO").
TRIM_KEEP_PARAGRAPH_COUNT = 96
COVER_TITLE_PARAGRAPH_INDEX = 15

DIAG_CAPAS = HERE / "_siif_diag_capas.png"
DIAG_HTTP = HERE / "_siif_diag_http.png"
DIAG_WORKER = HERE / "_siif_diag_worker.png"


def draw_architecture_diagram(path: Path) -> None:
    fig, ax = plt.subplots(1, 1, figsize=(9, 6.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 10)
    ax.axis("off")
    ax.set_title("Arquitectura por capas — SIIF Nación (resumen)", fontsize=14, fontweight="bold", pad=16)

    def box(x, y, w, h, text, fc="#E3F2FD", ec="#1565C0"):
        r = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.15",
            linewidth=1.5,
            edgecolor=ec,
            facecolor=fc,
        )
        ax.add_patch(r)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=8.5)

    box(2.8, 8.0, 4.4, 0.95, "API — ASP.NET Core 9\nControllers · Middleware · Swagger · Health")
    box(2.6, 6.2, 4.8, 1.05, "Application — MediatR, handlers, contratos")
    box(0.35, 3.9, 2.5, 1.25, "Domain\nentidades")
    box(3.75, 3.9, 2.5, 1.25, "Infrastructure\nSIIF HttpClient\ntoken · logging")
    box(7.15, 3.9, 2.5, 1.25, "Persistence\nEF Core · SQL Server\nrepositorios")

    def arrow(x0, y0, x1, y1):
        ax.add_patch(
            FancyArrowPatch(
                (x0, y0),
                (x1, y1),
                arrowstyle="-|>",
                mutation_scale=14,
                linewidth=1.2,
                color="#424242",
            )
        )

    arrow(5, 8.0, 5, 7.25)
    arrow(3.8, 6.2, 1.6, 5.15)
    arrow(5, 6.2, 5, 5.15)
    arrow(6.2, 6.2, 8.4, 5.15)
    ax.text(5.25, 7.5, "referencia", fontsize=7, color="#616161")
    ax.text(2.5, 5.9, "usa", fontsize=7, color="#616161")
    ax.text(5.15, 5.9, "interfaces →", fontsize=7, color="#616161", ha="right")
    ax.text(6.8, 5.9, "← interfaces", fontsize=7, color="#616161")

    plt.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_http_flow_diagram(path: Path) -> None:
    fig, ax = plt.subplots(1, 1, figsize=(10, 2.4))
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 2)
    ax.axis("off")
    ax.set_title("Flujo de una petición HTTP (resumen)", fontsize=13, fontweight="bold", pad=10)

    steps = [
        (0.2, 0.5, 1.6, 1.0, "HTTP\nPipeline"),
        (2.1, 0.5, 1.7, 1.0, "Controller\nDTO→Cmd/Qry"),
        (4.2, 0.5, 1.5, 1.0, "IMediator\nSend"),
        (6.0, 0.5, 1.6, 1.0, "Handler\n(Application)"),
        (7.9, 0.5, 1.8, 1.0, "ISiif*\n(HttpClient)"),
        (10.0, 0.5, 1.7, 1.0, "I*Repository\n(EF Core)"),
        (11.9, 0.5, 1.5, 1.0, "JSON\nIActionResult"),
    ]
    for x, y, w, h, t in steps:
        r = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.1",
            linewidth=1.2,
            edgecolor="#2E7D32",
            facecolor="#E8F5E9",
        )
        ax.add_patch(r)
        ax.text(x + w / 2, y + h / 2, t, ha="center", va="center", fontsize=7.5)

    for i in range(len(steps) - 1):
        x0 = steps[i][0] + steps[i][2]
        x1 = steps[i + 1][0]
        ymid = steps[i][1] + steps[i][3] / 2
        ax.add_patch(
            FancyArrowPatch(
                (x0 + 0.05, ymid),
                (x1 - 0.05, ymid),
                arrowstyle="-|>",
                mutation_scale=10,
                linewidth=1,
                color="#424242",
            )
        )

    plt.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_worker_diagram(path: Path) -> None:
    fig, ax = plt.subplots(1, 1, figsize=(10, 2.8))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 2.2)
    ax.axis("off")
    ax.set_title("Worker — ciclo y paralelismo (resumen)", fontsize=13, fontweight="bold", pad=10)

    steps = [
        (0.15, 0.55, 1.7, 1.05, "Cron +\nTimeZoneId"),
        (2.15, 0.55, 1.55, 1.05, "SiifScheduled\nIntegrationWorker"),
        (4.0, 0.55, 1.45, 1.05, "idCiclo\n(GUID)"),
        (5.75, 0.55, 1.55, 1.05, "Orquestador\n(scoped)"),
        (7.55, 0.55, 1.85, 1.05, "Paralelo:\nCDP · CDP pag · Oblig x3"),
        (9.65, 0.55, 1.55, 1.05, "Auditoría\nIAuditoriaLogger"),
    ]
    for x, y, w, h, t in steps:
        r = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.1",
            linewidth=1.2,
            edgecolor="#6A1B9A",
            facecolor="#F3E5F5",
        )
        ax.add_patch(r)
        ax.text(x + w / 2, y + h / 2, t, ha="center", va="center", fontsize=7.5)

    for i in range(len(steps) - 1):
        x0 = steps[i][0] + steps[i][2]
        x1 = steps[i + 1][0]
        ymid = steps[i][1] + steps[i][3] / 2
        ax.add_patch(
            FancyArrowPatch(
                (x0 + 0.05, ymid),
                (x1 - 0.05, ymid),
                arrowstyle="-|>",
                mutation_scale=10,
                linewidth=1,
                color="#424242",
            )
        )

    plt.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def trim_template_after_cover(doc: Document, keep_paragraphs: int) -> None:
    """Elimina el contenido del cuerpo posterior al párrafo (keep_paragraphs-1), conservando portada y sectPr."""
    if keep_paragraphs < 1:
        raise ValueError("keep_paragraphs debe ser >= 1")
    last_idx = keep_paragraphs - 1
    if last_idx >= len(doc.paragraphs):
        return
    anchor = doc.paragraphs[last_idx]._element
    parent = anchor.getparent()
    sib = anchor.getnext()
    while sib is not None:
        tag = sib.tag
        nxt = sib.getnext()
        if tag.endswith("sectPr"):
            break
        parent.remove(sib)
        sib = nxt


def set_cover_title(doc: Document, title: str) -> None:
    if COVER_TITLE_PARAGRAPH_INDEX < len(doc.paragraphs):
        doc.paragraphs[COVER_TITLE_PARAGRAPH_INDEX].text = title


def _clean_inline(text: str) -> str:
    t = text.replace("\u00a0", " ").strip()
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"`([^`]+)`", r"\1", t)
    return t


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows: list[list[str]] = []
    for tl in table_lines:
        if re.match(r"^\|\s*:?-{3,}", tl.strip()):
            continue
        raw = [c.strip() for c in tl.strip().strip("|").split("|")]
        rows.append(raw)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    try:
        tbl.style = "Table Grid"
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell_text = row[ci] if ci < len(row) else ""
            tbl.rows[ri].cells[ci].text = _clean_inline(cell_text)


def append_markdown(doc: Document, md_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    first_h1_skipped = False

    def flush_code() -> None:
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_buf))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                i += 1
                continue
            flush_code()
            in_code = False
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        stripped = line.strip()

        if stripped.startswith("# ") and not first_h1_skipped:
            first_h1_skipped = True
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("|") and stripped.count("|") >= 2:
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_markdown_table(doc, block)
            continue

        if stripped.startswith("#### "):
            doc.add_heading(_clean_inline(stripped[5:]), level=3)
        elif stripped.startswith("### "):
            doc.add_heading(_clean_inline(stripped[4:]), level=2)
        elif stripped.startswith("## "):
            doc.add_heading(_clean_inline(stripped[3:]), level=1)
        elif stripped.startswith("# "):
            doc.add_heading(_clean_inline(stripped[2:]), level=1)
        elif re.match(r"^\d+\.\s+", stripped):
            body = re.sub(r"^\d+\.\s+", "", stripped)
            try:
                doc.add_paragraph(_clean_inline(body), style="List Number")
            except (KeyError, ValueError):
                doc.add_paragraph(_clean_inline(body))
        elif stripped.startswith("- "):
            try:
                doc.add_paragraph(_clean_inline(stripped[2:]), style="List Bullet")
            except (KeyError, ValueError):
                doc.add_paragraph("• " + _clean_inline(stripped[2:]))
        elif stripped:
            doc.add_paragraph(_clean_inline(stripped))
        else:
            doc.add_paragraph("")

        i += 1

    if in_code:
        flush_code()


def insert_figure(doc: Document, png_path: Path, caption: str) -> None:
    if not png_path.is_file():
        return
    cap = doc.add_paragraph()
    cap_run = cap.add_run(caption)
    cap_run.italic = True
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(str(png_path), width=Inches(6.2))


def build_architecture_docx(out_path: Path) -> None:
    draw_architecture_diagram(DIAG_CAPAS)
    draw_http_flow_diagram(DIAG_HTTP)

    doc = Document(str(TEMPLATE_DOCX))
    trim_template_after_cover(doc, TRIM_KEEP_PARAGRAPH_COUNT)
    set_cover_title(doc, "Documento de Arquitectura Interoperabilidad SIIF Nación SSF.")

    append_markdown(doc, HERE / "Arquitectura_SIIFNacion.md")

    doc.add_heading("Figuras", level=1)
    insert_figure(doc, DIAG_CAPAS, "Figura 1. Arquitectura por capas (resumen).")
    insert_figure(doc, DIAG_HTTP, "Figura 2. Flujo de una petición HTTP (resumen).")

    doc.save(str(out_path))


def build_manual_docx(
    md_name: str,
    out_name: str,
    cover_title: str,
    *,
    worker_diagram: bool = False,
) -> None:
    if worker_diagram:
        draw_worker_diagram(DIAG_WORKER)

    doc = Document(str(TEMPLATE_DOCX))
    trim_template_after_cover(doc, TRIM_KEEP_PARAGRAPH_COUNT)
    set_cover_title(doc, cover_title)

    append_markdown(doc, HERE / md_name)

    if worker_diagram and DIAG_WORKER.is_file():
        doc.add_heading("Figura — ciclo del Worker", level=1)
        insert_figure(doc, DIAG_WORKER, "Figura 1. Worker: programación, ciclo y paralelismo (resumen).")

    doc.save(str(HERE / out_name))


def main() -> None:
    if not TEMPLATE_DOCX.is_file():
        raise FileNotFoundError(f"No se encontró la plantilla: {TEMPLATE_DOCX}")

    build_architecture_docx(HERE / "Arquitectura_SIIFNacion.docx")

    build_manual_docx(
        "Manual_Instalador.md",
        "Manual_Instalador_SIIFNacion.docx",
        "Manual de uso del instalador — SIIF Nación SSF.",
    )
    build_manual_docx(
        "Manual_APIs.md",
        "Manual_APIs_SIIFNacion.docx",
        "Manual de uso de las APIs REST — SIIF Nación SSF.",
    )
    build_manual_docx(
        "Manual_Worker.md",
        "Manual_Worker_SIIFNacion.docx",
        "Manual de uso del Worker — SIIF Nación SSF.",
        worker_diagram=True,
    )

    print("Generados en", HERE)
    for name in (
        "Arquitectura_SIIFNacion.docx",
        "Manual_Instalador_SIIFNacion.docx",
        "Manual_APIs_SIIFNacion.docx",
        "Manual_Worker_SIIFNacion.docx",
    ):
        p = HERE / name
        print(" ", p.name, f"({p.stat().st_size // 1024} KB)" if p.is_file() else "ERROR")


if __name__ == "__main__":
    main()
