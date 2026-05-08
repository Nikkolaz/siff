# -*- coding: utf-8 -*-
"""Genera VISION-GENERAL-PROYECTO.docx con diagramas (PNG) y el contenido de visión del proyecto."""
from __future__ import annotations

import os
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
OUT_DOCX = HERE / "VISION-GENERAL-PROYECTO.docx"
DIAG1 = HERE / "_diagrama_arquitectura_capas.png"
DIAG2 = HERE / "_diagrama_flujo_peticion.png"


def draw_architecture_diagram(path: Path) -> None:
    fig, ax = plt.subplots(1, 1, figsize=(9, 6.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 10)
    ax.axis("off")
    ax.set_title("Arquitectura por capas (resumen)", fontsize=14, fontweight="bold", pad=16)

    def box(x, y, w, h, text, fc="#E3F2FD", ec="#1565C0"):
        r = FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.15",
            linewidth=1.5, edgecolor=ec, facecolor=fc,
        )
        ax.add_patch(r)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=8.5)

    box(2.8, 8.0, 4.4, 0.95, "API — ASP.NET Core 9\nControllers · Middleware · Swagger · Health")
    box(2.6, 6.2, 4.8, 1.05, "Application — MediatR, handlers, DTOs, contratos")
    box(0.35, 3.9, 2.5, 1.25, "Domain\nentidades\n(legado Service.*)")
    box(3.75, 3.9, 2.5, 1.25, "Infrastructure\nSIIF HttpClient\ntoken · logging")
    box(7.15, 3.9, 2.5, 1.25, "Persistence\nEF Core · SQL Server\nrepositorios")

    def arrow(x0, y0, x1, y1):
        ax.add_patch(
            FancyArrowPatch(
                (x0, y0), (x1, y1),
                arrowstyle="-|>", mutation_scale=14, linewidth=1.2, color="#424242",
            )
        )

    arrow(5, 8.0, 5, 7.25)  # API → Application
    arrow(3.8, 6.2, 1.6, 5.15)  # App → Domain
    arrow(5, 6.2, 5, 5.15)  # App → Infra
    arrow(6.2, 6.2, 8.4, 5.15)  # App → Persistence
    ax.text(5.25, 7.5, "referencia", fontsize=7, color="#616161")
    ax.text(2.5, 5.9, "usa", fontsize=7, color="#616161")
    ax.text(5.15, 5.9, "interfaces →", fontsize=7, color="#616161", ha="right")
    ax.text(6.8, 5.9, "← interfaces", fontsize=7, color="#616161")

    plt.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_request_flow_diagram(path: Path) -> None:
    fig, ax = plt.subplots(1, 1, figsize=(10, 2.4))
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 2)
    ax.axis("off")
    ax.set_title("Flujo de una petición HTTP (resumen)", fontsize=13, fontweight="bold", pad=10)

    steps = [
        (0.2, 0.5, 1.6, 1.0, "HTTP\nPipeline"),
        (2.1, 0.5, 1.7, 1.0, "Controller\nDTO→Cmd/Qry"),
        (4.2, 0.5, 1.5, 1.0, "IMediator\nSend"),
        (6.0, 0.5, 1.6, 1.0, "Handler\n(Application)"),
        (7.9, 0.5, 1.8, 1.0, "ISiif*\n(HttpClient)"),
        (10.0, 0.5, 1.7, 1.0, "I*Repository\n(EF Core)"),
        (11.9, 0.5, 1.5, 1.0, "JSON\nIActionResult"),
    ]
    for x, y, w, h, t in steps:
        r = FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.1",
            linewidth=1.2, edgecolor="#2E7D32", facecolor="#E8F5E9",
        )
        ax.add_patch(r)
        ax.text(x + w / 2, y + h / 2, t, ha="center", va="center", fontsize=7.5)

    for i in range(len(steps) - 1):
        x0 = steps[i][0] + steps[i][2]
        x1 = steps[i + 1][0]
        ymid = steps[i][1] + steps[i][3] / 2
        ax.add_patch(
            FancyArrowPatch(
                (x0 + 0.05, ymid), (x1 - 0.05, ymid),
                arrowstyle="-|>", mutation_scale=10, linewidth=1, color="#424242",
            )
        )

    plt.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_mermaid_block(doc: Document, title: str, mermaid_src: str) -> None:
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    run = p.add_run(
        "Puede copiar el siguiente bloque en https://mermaid.live para editar o exportar como imagen."
    )
    run.italic = True
    run.font.size = Pt(9)
    code_p = doc.add_paragraph()
    code_run = code_p.add_run(mermaid_src)
    code_run.font.name = "Consolas"
    code_run.font.size = Pt(7)


def build_document() -> None:
    draw_architecture_diagram(DIAG1)
    draw_request_flow_diagram(DIAG2)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    t = doc.add_heading("Visión general — SSF.Interop.SIIFNacion.API", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Documento de síntesis para onboarding y contexto antes de modificar el código. "
        "Stack: ASP.NET Core 9, MediatR, EF Core (SQL Server), integración HTTP con SIIF (MinHacienda)."
    )

    doc.add_heading("Diagramas", level=1)
    doc.add_paragraph("Figura 1 — Arquitectura por capas (vista resumida).")
    doc.add_picture(str(DIAG1), width=Inches(6.2))
    doc.add_paragraph()
    doc.add_paragraph("Figura 2 — Flujo de una petición HTTP.")
    doc.add_picture(str(DIAG2), width=Inches(6.5))

    mermaid_arch = """flowchart TB
  subgraph host [API - ASP.NET Core 9]
    Program[Program.cs]
    Ctrl[Controllers]
    MW[ExceptionMiddleware]
    Program --> Ctrl
    Program --> MW
  end
  subgraph app [Application]
    Med[MediatR Handlers]
    DTO[DTOs / Queries / Commands]
    Contracts[Contracts e interfaces SIIF y Persistence]
    Med --> Contracts
  end
  subgraph dom [Domain]
    Ent[Entidades CDP, SIIF, Auditoría, Queue stubs]
  end
  subgraph infra [Infrastructure]
    Http[HttpClient SIIF + DelegatingHandlers]
    Tok[Token + caché memoria]
    Log[Logging compuesto]
    Net[WebClientWrapper]
    Http --> Tok
  end
  subgraph pers [Persistence]
    EF[EF Core SQL Server]
    Repo[Repositorios CDP / Obligaciones]
  end
  Ctrl -->|IMediator| Med
  Med -->|ISiif*| Http
  Med -->|I*Repository| Repo
  app --> dom
  infra --> Contracts
  pers --> Contracts
  host --> app
  host --> infra
  host --> pers"""

    add_mermaid_block(doc, "Diagrama Mermaid — arquitectura detallada", mermaid_arch)

    # --- Secciones del markdown (sintetizadas) ---
    sections = [
        (
            "1. Qué es esta aplicación",
            [
                "API REST de interoperabilidad con SIIF Nación: expone endpoints que orquestan llamadas al sistema SIIF "
                "y persisten resultados en SQL Server (GESTORDOC / tablas relacionadas). Casos de uso con MediatR; "
                "estructura de capas tipo Clean Architecture ligera.",
            ],
        ),
        (
            "2. Proyectos de la solución",
            [
                "SSF.Interop.SIIFNacion.API — Host web: controladores, middleware, Swagger/OpenAPI, health checks.",
                "SSF.Interop.SIIFNacion.Application — MediatR, DTOs, contratos, AutoMapper, FluentValidation. Depende de Domain.",
                "SSF.Interop.SIIFNacion.Domain — Entidades (CDP, obligaciones SIIF, auditoría/cola heredada).",
                "SSF.Interop.SIIFNacion.Infrastructure — Clientes HTTP SIIF, opciones, token/caché, logging, WebClientWrapper, storage local.",
                "SSF.Interop.SIIFNacion.Persistence — EF Core, DbContext, repositorios.",
                "En el .sln: carpetas src → API / Core / Infrastructure. Carpeta Test sin proyecto de pruebas referenciado.",
                "Dependencias: API → Application, Infrastructure, Persistence | Application → Domain | Infrastructure → Application | Persistence → Application, Domain.",
            ],
        ),
        (
            "3. Punto de entrada y arranque",
            [
                "Archivo: SSF.Interop.SIIFNacion.API/Program.cs",
                "Flujo: WebApplication.CreateBuilder → ConfigureServices → Build → pipeline HTTP → Run().",
                "Extensiones: ConfigureApplicationServices (MediatR, AutoMapper); ConfigureInfrastructureServices (SIIF, logging); "
                "ConfigurePresistenceServices (auditoría + proceso) — los contextos EF usan hoy la cadena de proceso; auditoría puede quedar sin uso en ese registro.",
                "Pipeline aproximado: HTTPS, UseAuthorization, MapControllers, Swagger, ExceptionMiddleware, health checks.",
                "Notas: AddControllers duplicado; CORS registrado — verificar UseCors en pipeline si se necesita.",
            ],
        ),
        (
            "4. Arquitectura por capas / módulos",
            [
                "API: HTTP ↔ Commands/Queries; validación mínima en controladores.",
                "Application: handlers MediatR; solo interfaces (ISiif*, I*Repository).",
                "Domain: entidades; parte del código en namespaces SSF.Interop.SIIFNacion.Service.* (legado / copia de otro servicio).",
                "Infrastructure: SIIF (HttpClient, Bearer, retry 401), token en memoria, logging.",
                "Persistence: EF Core + SQL Server.",
            ],
        ),
        (
            "5. Acceso a datos",
            [
                "EF Core + SqlServer. Contextos: GestordocDbContext, DbContextCdp.",
                "Repositorios: ICdpRepository, ICdpPaginadoRepository, IObligacionApoRepository.",
                "ConnectionStrings en appsettings; Program también referencia SqlServerAuditoria al armar el par de cadenas.",
            ],
        ),
        (
            "6. Autenticación y seguridad",
            [
                "Hacia consumidores de la API: no hay esquema típico JWT/[Authorize] evidente; launchSettings sugiere anónimo.",
                "Hacia SIIF: Siif:User/Password → SiifAuthenticationService → SiifTokenProvider (caché) → SiifAuthDelegatingHandler (Bearer) → SiifRetryOnUnauthorizedHandler.",
                "Riesgo: secretos en JSON; externalizar (User Secrets, env, AWS).",
            ],
        ),
        (
            "7. Integraciones externas",
            [
                "SIIF (MinHacienda): BaseUrl + ISiifAuthenticationService, ISiifBudgetService, ISiifCatalogService.",
                "WebClientWrapper: HttpClient genérico (distinto a IHttpClientFactory de SIIF).",
                "AWS SDK en Infrastructure — confirmar uso activo.",
            ],
        ),
        (
            "8. Configuración crítica",
            [
                "appsettings.json, Development, Qa: Environment, Siif (BaseUrl, AuthenticationPath, User, Password, tokens, Timeout), "
                "ConnectionStrings (SqlServerProceso, SqlServerAuditoria en código), WorkerSettings:Convenio.",
                "SiifOptions.cs — SectionName = Siif.",
            ],
        ),
        (
            "9. Flujos de negocio principales",
            [
                "POST api/siif/consultar-cdp → ConsultarCdpQuery → SIIF + persistencia según respuesta.",
                "POST api/siif/consultar-cdp-paginada → ConsultarCdpPaginadaQuery + repositorio paginado.",
                "POST api/siif/sincronizar-lista-obligaciones → comando MediatR → paginación SIIF (máx. 50/página) → DYNTBLSIIFOBLIGAPO.",
                "Health: /healthcheck, UI; GET / en DefaultController.",
            ],
        ),
        (
            "10. Archivos clave",
            [
                "Program.cs; ApplicationServiceRegistration; InfrastructureServicesRegistration; PersistenceServicesRegistration.",
                "SiifController, DefaultController, ExceptionMiddleware.",
                "Handlers: ConsultarCdpQueryHandler, ConsultarCdpPaginadaQueryHandler, SincronizarListaObligacionesCommandHandler.",
                "SIIF: SiifAuthenticationService, SiifTokenProvider, SiifAuthDelegatingHandler, SiifRetryOnUnauthorizedHandler, SiifBudgetService, SiifOptions.",
                "Persistence: GestordocDbContext, CdpRepository, CdpPaginadoRepository, ObligacionApoRepository.",
                "Contratos: Application/Common/Interfaces/ExternalServices/SIIF/*, Contracts/Persistence ICdp*, IObligacionApoRepository.",
            ],
        ),
        (
            "11. Riesgos y legado",
            [
                "Namespaces Service.* mezclados con SIIFNacion.*.",
                "Contratos no registrados (auditoría, IUnitOfWork, genéricos) — posible código muerto.",
                "WebClientWrapper grande; HttpClient sin factory en parte del código.",
                "Handlers extensos (sincronizar obligaciones).",
                "SemaphoreSlim estático en SiifTokenProvider.",
                "DatabaseLogger + IEnumerable<ICustomLogger> + CompositeLogger: riesgo de ciclo/recursión.",
                "ExceptionMiddleware después de MapControllers: puede no capturar excepciones de controladores como se espera.",
                "Secretos en configuración; AddControllers duplicado; CORS sin UseCors verificado.",
            ],
        ),
        (
            "12. Lista de lectura priorizada",
            [
                "Nivel 1: Program.cs, *ServiceRegistration.cs (Application, Infrastructure, Persistence), SiifController, appsettings.*",
                "Nivel 2: tres handlers SIIF, SiifBudgetService, autenticación/token/handlers HTTP, SiifOptions, GestordocDbContext + repos.",
                "Nivel 3: ExceptionMiddleware, ISiif*, WebClientWrapper, loggers, Domain/Auditoria y contratos Service.Application.Contracts.Persistence.",
            ],
        ),
        (
            "13. Flujo petición (texto)",
            [
                "HTTP → Pipeline (HTTPS, Auth ASP.NET, Routing) → Controller (DTO → Command/Query) → IMediator.Send → "
                "Handler → ISiif* (Infrastructure) / I*Repository (Persistence) → IActionResult JSON.",
            ],
        ),
    ]

    doc.add_page_break()
    doc.add_heading("Contenido detallado", level=1)

    for heading, paras in sections:
        doc.add_heading(heading, level=2)
        for line in paras:
            doc.add_paragraph(line, style=None)

    doc.add_paragraph()
    foot = doc.add_paragraph(
        "Documento generado automáticamente a partir de docs/VISION-GENERAL-PROYECTO.md y el análisis de arquitectura. "
        "Actualizar si cambian DI, endpoints o integraciones."
    )
    foot.runs[0].font.size = Pt(9)
    foot.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.save(OUT_DOCX)

    # Limpieza PNG temporales (opcional: mantener para re-export)
    for p in (DIAG1, DIAG2):
        try:
            os.remove(p)
        except OSError:
            pass


if __name__ == "__main__":
    build_document()
    print(f"Escrito: {OUT_DOCX}")
